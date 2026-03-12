"""Tests for brand_linter.parser.

Since we can't depend on real .docx fixtures being present in CI, we build
minimal in-memory documents with python-docx and write them to a tmp_path,
then parse them back.  This exercises the full round-trip through python-docx's
file I/O and the parser's resolution logic.
"""

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt, RGBColor

from brand_linter.parser import DocumentParseError, parse_document
from brand_linter.models import ParagraphData


# ---------------------------------------------------------------------------
# Fixture helpers
# ---------------------------------------------------------------------------


def make_doc(tmp_path: Path, paragraphs: list[dict]) -> Path:
    """Create a minimal .docx with the requested paragraph specs.

    Each dict in *paragraphs* may contain:
        text       : str
        style      : str   (Word style name, default "Normal")
        font_name  : str
        font_size  : float (points)
        bold       : bool
        italic     : bool
        color      : tuple(r, g, b)  as ints 0-255
    """
    doc = Document()
    # Remove the default blank paragraph Word always creates
    for p in doc.paragraphs:
        p._element.getparent().remove(p._element)

    for spec in paragraphs:
        style = spec.get("style", "Normal")
        text = spec.get("text", "")
        try:
            para = doc.add_paragraph(text, style=style)
        except KeyError:
            # Style doesn't exist in the blank doc template; fall back
            para = doc.add_paragraph(text)

        if para.runs:
            run = para.runs[0]
        else:
            run = para.add_run()

        if "font_name" in spec:
            run.font.name = spec["font_name"]
        if "font_size" in spec:
            run.font.size = Pt(spec["font_size"])
        if "bold" in spec:
            run.font.bold = spec["bold"]
        if "italic" in spec:
            run.font.italic = spec["italic"]
        if "color" in spec:
            r, g, b = spec["color"]
            run.font.color.rgb = RGBColor(r, g, b)

    path = tmp_path / "test.docx"
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Happy-path tests
# ---------------------------------------------------------------------------


def test_basic_paragraph_count(tmp_path):
    path = make_doc(tmp_path, [
        {"text": "Hello world"},
        {"text": "Second paragraph"},
    ])
    paras = parse_document(path)
    assert len(paras) == 2


def test_paragraph_index_and_text(tmp_path):
    path = make_doc(tmp_path, [
        {"text": "First"},
        {"text": "Second"},
        {"text": "Third"},
    ])
    paras = parse_document(path)
    assert paras[0].index == 0
    assert paras[0].text == "First"
    assert paras[2].index == 2
    assert paras[2].text == "Third"


def test_font_name_captured(tmp_path):
    path = make_doc(tmp_path, [{"text": "Styled", "font_name": "Arial"}])
    paras = parse_document(path)
    assert paras[0].font_name == "Arial"


def test_font_size_captured_as_float(tmp_path):
    path = make_doc(tmp_path, [{"text": "Sized", "font_size": 14}])
    paras = parse_document(path)
    assert paras[0].font_size == pytest.approx(14.0)


def test_bold_captured(tmp_path):
    path = make_doc(tmp_path, [{"text": "Bold text", "bold": True}])
    paras = parse_document(path)
    assert paras[0].bold is True


def test_italic_captured(tmp_path):
    path = make_doc(tmp_path, [{"text": "Italic text", "italic": True}])
    paras = parse_document(path)
    assert paras[0].italic is True


def test_color_captured_as_hex_string(tmp_path):
    path = make_doc(tmp_path, [{"text": "Colored", "color": (0x1F, 0x38, 0x64)}])
    paras = parse_document(path)
    # RGBColor.__str__ returns uppercase hex
    assert paras[0].color_hex is not None
    assert paras[0].color_hex.upper() == "1F3864"


def test_empty_paragraph(tmp_path):
    path = make_doc(tmp_path, [{"text": ""}])
    paras = parse_document(path)
    assert len(paras) == 1
    assert paras[0].text == ""


def test_returns_paragraph_data_instances(tmp_path):
    path = make_doc(tmp_path, [{"text": "x"}])
    paras = parse_document(path)
    assert all(isinstance(p, ParagraphData) for p in paras)


def test_ref_format(tmp_path):
    path = make_doc(tmp_path, [{"text": "a"}, {"text": "b"}])
    paras = parse_document(path)
    assert paras[0].ref == "para[0]"
    assert paras[1].ref == "para[1]"


# ---------------------------------------------------------------------------
# Error-path tests
# ---------------------------------------------------------------------------


def test_missing_file_raises():
    with pytest.raises(FileNotFoundError):
        parse_document("/nonexistent/doc.docx")


def test_wrong_extension_raises(tmp_path):
    p = tmp_path / "doc.txt"
    p.write_text("hello")
    with pytest.raises(DocumentParseError, match=".docx"):
        parse_document(p)
